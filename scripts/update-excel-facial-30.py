#!/usr/bin/env python3
"""
Actualiza las 30 filas de Cuidado Facial en la plantilla de Google Drive:
Mesolabpro_Plantilla_Seleccion_Productos.xlsx (filas 32 a 61)
"""

import openpyxl

EXCEL_PATH = "/Users/cristianborrero/Library/CloudStorage/GoogleDrive-infomesolabpro@gmail.com/My Drive/Mesolabpro/06_Catalogo_y_Productos/02_Estudios_de_Mercado_y_Ganadores/Mesolabpro_Plantilla_Seleccion_Productos.xlsx"

FACIAL_PRODUCTS = [
    {
        "id": "031", "cat": "Cuidado Facial", "sub": "Sueros Dermo-Activos",
        "name": "Suero Hidratante Concentrado Ácido Hialurónico Puro 30ml",
        "supplier": "345873", "dropiId": "1607896", "costo": 26000, "pvp": 69000, "stock": 592,
        "slug": "suero-acido-hialuronico-puro-30ml", "nota": "Ácido hialurónico multimolecular, alta retención hídrica."
    },
    {
        "id": "032", "cat": "Cuidado Facial", "sub": "Sueros Dermo-Activos",
        "name": "Suero Iluminador Antioxidante Vitamina C + Vitamina E",
        "supplier": "7910", "dropiId": "2239906", "costo": 18000, "pvp": 59000, "stock": 400,
        "slug": "suero-iluminador-vitamina-c-vitamina-e", "nota": "Complejo antioxidante que unifica el tono y combate radicales libres."
    },
    {
        "id": "033", "cat": "Cuidado Facial", "sub": "Tratamientos Antiedad y Regeneradores",
        "name": "Suero Reparador Nocturno Retinol Puro Antienvejecimiento",
        "supplier": "43728", "dropiId": "2117094", "costo": 13000, "pvp": 49000, "stock": 114,
        "slug": "suero-reparador-retinol-antienvejecimiento", "nota": "Retinol encapsulado con centella asiática para renovación sin irritación."
    },
    {
        "id": "034", "cat": "Cuidado Facial", "sub": "Sueros Dermo-Activos",
        "name": "Suero Niacinamida 10% + Zinc Regulador de Sebo y Poros",
        "supplier": "839", "dropiId": "2230851", "costo": 15000, "pvp": 52000, "stock": 996,
        "slug": "suero-niacinamida-10-zinc-seborregulador", "nota": "Niacinamida 10% + Zinc PCA 1% para control de brillo y minimización de poros."
    },
    {
        "id": "035", "cat": "Cuidado Facial", "sub": "Tónicos y Brumas Faciales",
        "name": "Tónico Exfoliante con Ácido Glicólico AHA al 7%",
        "supplier": "81817", "dropiId": "2151199", "costo": 13500, "pvp": 49000, "stock": 4993,
        "slug": "tonico-exfoliante-acido-glicolico-7", "nota": "AHA al 7% para exfoliación química suave, textura sedosa y luminosidad."
    },
    {
        "id": "036", "cat": "Cuidado Facial", "sub": "Tratamientos Antiacné y Control Grasa",
        "name": "Suero Purificante Ácido Salicílico BHA 2% Antiacné",
        "supplier": "120432", "dropiId": "1749325", "costo": 19200, "pvp": 59000, "stock": 259,
        "slug": "suero-acido-salicilico-bha-2-antiacne", "nota": "BHA 2% lipofílico que disuelve grasa y comedones dentro del poro."
    },
    {
        "id": "037", "cat": "Cuidado Facial", "sub": "Protección Solar y Filtros UV",
        "name": "Protector Solar Toque Seco Colágeno FPS 60 UVA/UVB",
        "supplier": "144205", "dropiId": "1849339", "costo": 10000, "pvp": 45000, "stock": 145,
        "slug": "protector-solar-toque-seco-colageno-fps60", "nota": "FPS 60 no graso con colágeno hidrolizado sin efecto blanco."
    },
    {
        "id": "038", "cat": "Cuidado Facial", "sub": "Mascarillas y Velo de Colágeno",
        "name": "Mascarilla Facial Bio-Colágeno Hidrogel Profundo (Pack x5)",
        "supplier": "76117", "dropiId": "2259141", "costo": 10500, "pvp": 45000, "stock": 112,
        "slug": "mascarilla-facial-biocolageno-hidrogel-pack5", "nota": "Láminas de hidrogel transdérmico para hidratación oclusiva intensiva."
    },
    {
        "id": "039", "cat": "Cuidado Facial", "sub": "Cuidado de Ojos y Ojeras",
        "name": "Contorno de Ojos Roll-On Drenante con Ácido Hialurónico",
        "supplier": "899060", "dropiId": "2207878", "costo": 8000, "pvp": 39000, "stock": 1700,
        "slug": "contorno-ojos-roll-on-acido-hialuronico", "nota": "Triple esfera metálica que proporciona masaje frío drenante contra bolsas."
    },
    {
        "id": "040", "cat": "Cuidado Facial", "sub": "Limpieza y Desmaquillantes",
        "name": "Espuma Limpiadora Facial con Aminoácidos y Cepillo 200ml",
        "supplier": "345873", "dropiId": "1637125", "costo": 16000, "pvp": 49000, "stock": 496,
        "slug": "espuma-limpiadora-aminoacidos-cepillo-200ml", "nota": "Surfactantes suaves de aminoácidos con cabezal de silicona integrado."
    },
    {
        "id": "041", "cat": "Cuidado Facial", "sub": "Sueros Dermo-Activos",
        "name": "Suero Facial Regenerador de Centella Asiática (Cica) 40ml",
        "supplier": "43339", "dropiId": "2183092", "costo": 19970, "pvp": 59000, "stock": 4000,
        "slug": "suero-facial-centella-asiatica-cica-40ml", "nota": "Extracto calmante y reparador de barrera dérmica para pieles reactivas."
    },
    {
        "id": "042", "cat": "Cuidado Facial", "sub": "Tratamientos Antiedad y Regeneradores",
        "name": "Crema Reafirmante Antiarrugas con Caviar y Péptidos Tensores",
        "supplier": "45331", "dropiId": "2037092", "costo": 17500, "pvp": 59000, "stock": 142,
        "slug": "crema-antiarrugas-caviar-peptidos-tensores", "nota": "Extracto de caviar negro y hexapéptidos tensores para densidad dérmica."
    },
    {
        "id": "043", "cat": "Cuidado Facial", "sub": "Limpieza y Desmaquillantes",
        "name": "Gel Limpiador Facial Purificante con Extracto de Baba de Caracol",
        "supplier": "43339", "dropiId": "2187531", "costo": 8970, "pvp": 39000, "stock": 444,
        "slug": "limpiador-facial-baba-caracol-purificante", "nota": "Mucina de caracol filtrada que regenera y limpia sin resecar."
    },
    {
        "id": "044", "cat": "Cuidado Facial", "sub": "Tónicos y Brumas Faciales",
        "name": "Tónico Astringente Facial de Caléndula y Manzanilla",
        "supplier": "440677", "dropiId": "1726680", "costo": 34000, "pvp": 79000, "stock": 162,
        "slug": "tonico-facial-calendula-manzanilla-astringente", "nota": "Pétalos de caléndula infusionados para calmar rojeces y equilibrar pH."
    },
    {
        "id": "045", "cat": "Cuidado Facial", "sub": "Limpieza y Desmaquillantes",
        "name": "Desmaquillante Micelar Bifásico Purificante 150ml",
        "supplier": "184977", "dropiId": "915481", "costo": 11888, "pvp": 39000, "stock": 1998,
        "slug": "desmaquillante-micelar-bifasico-purificante-150ml", "nota": "Fase oleosa y acuosa micelar que remueve maquillaje a prueba de agua."
    },
    {
        "id": "046", "cat": "Cuidado Facial", "sub": "Tónicos y Brumas Faciales",
        "name": "Agua de Rosas Orgánica Tonificante y Equilibrante 250ml",
        "supplier": "184977", "dropiId": "915469", "costo": 6240, "pvp": 29000, "stock": 5624,
        "slug": "agua-de-rosas-organica-tonificante-250ml", "nota": "Hidrolato puro 100% natural libre de alcohol con aplicador en spray."
    },
    {
        "id": "047", "cat": "Cuidado Facial", "sub": "Exfoliantes y Peeling Químico",
        "name": "Gel Exfoliante Dermo-Activo con Microesferas y Ácido Hialurónico",
        "supplier": "184977", "dropiId": "915482", "costo": 15600, "pvp": 49000, "stock": 2155,
        "slug": "gel-exfoliante-dermoactivo-acido-hialuronico", "nota": "Gommage enzimático suave que desprende células muertas por fricción."
    },
    {
        "id": "048", "cat": "Cuidado Facial", "sub": "Tratamientos Antiedad y Regeneradores",
        "name": "Crema Hidratante Reafirmante con Retinol y Colágeno 120g",
        "supplier": "596479", "dropiId": "2207193", "costo": 10000, "pvp": 45000, "stock": 9909,
        "slug": "crema-hidratante-retinol-colageno-120g", "nota": "Nutrición antiedad nocturna para rostro, cuello y escote."
    },
    {
        "id": "049", "cat": "Cuidado Facial", "sub": "Tratamientos Aclarantes y Despigmentantes",
        "name": "Crema Aclarante Antimanchas con Vitamina C y Niacinamida",
        "supplier": "29151", "dropiId": "1224273", "costo": 16100, "pvp": 49000, "stock": 498,
        "slug": "crema-aclarante-vitamina-c-niacinamida", "nota": "Corrección cromática contra manchas solares y secuelas de acné."
    },
    {
        "id": "050", "cat": "Cuidado Facial", "sub": "Hidratantes y Nutrición Celular",
        "name": "Gel Hidratante Ultrafresco con Ácido Hialurónico Bioaqua",
        "supplier": "899060", "dropiId": "2150711", "costo": 7000, "pvp": 35000, "stock": 190,
        "slug": "gel-hidratante-acido-hialuronico-bioaqua", "nota": "Textura acuosa Oil-Free para hidratación ligera en climas cálidos."
    },
    {
        "id": "051", "cat": "Cuidado Facial", "sub": "Cuidado de Ojos y Ojeras",
        "name": "Crema Contorno de Ojos Despigmentante Antimanchas con Péptidos",
        "supplier": "899060", "dropiId": "2150696", "costo": 9000, "pvp": 39000, "stock": 495,
        "slug": "crema-contorno-ojos-despigmentante-peptidos", "nota": "Tratamiento periocular con cafeína y péptidos aclarantes."
    },
    {
        "id": "052", "cat": "Cuidado Facial", "sub": "Cuidado de Ojos y Ojeras",
        "name": "Contorno de Ojos Reafirmante con Centella Asiática Bioaqua",
        "supplier": "23739", "dropiId": "1183762", "costo": 9000, "pvp": 39000, "stock": 988,
        "slug": "contorno-ojos-centella-asiatica-bioaqua", "nota": "Alivio calmante periocular que reduce hinchazón y patas de gallo."
    },
    {
        "id": "053", "cat": "Cuidado Facial", "sub": "Tratamientos Antiacné y Control Grasa",
        "name": "Gel Concentrado Anti-Imperfecciones Ácido Salicílico Bioaqua",
        "supplier": "39425", "dropiId": "217021", "costo": 15000, "pvp": 45000, "stock": 1874,
        "slug": "gel-anti-imperfecciones-acido-salicilico-bioaqua", "nota": "Acción secante focalizada sobre granos activos en 24 horas."
    },
    {
        "id": "054", "cat": "Cuidado Facial", "sub": "Sueros Dermo-Activos",
        "name": "Suero Facial Glow Nutritivo con Própolis y Niacinamida 30ml",
        "supplier": "19662", "dropiId": "2176538", "costo": 15000, "pvp": 52000, "stock": 500,
        "slug": "suero-facial-glow-propolis-niacinamida", "nota": "Inspirado en K-Beauty para efecto piel de cristal (glass skin) antibacteriano."
    },
    {
        "id": "055", "cat": "Cuidado Facial", "sub": "Protección Solar y Filtros UV",
        "name": "Bloqueador Solar Facial Mineral con Filtros Físicos SPF 50+",
        "supplier": "383592", "dropiId": "2176643", "costo": 17550, "pvp": 55000, "stock": 3638,
        "slug": "bloqueador-solar-facial-mineral-spf50", "nota": "Óxido de zinc y titanio que reflejan la radiación UV de forma mineral."
    },
    {
        "id": "056", "cat": "Cuidado Facial", "sub": "Kits y Rutinas Completas",
        "name": "Kit Completo Dermo-Regenerador Vitamina C Bioaqua (5 Pasos)",
        "supplier": "959749", "dropiId": "2227720", "costo": 33000, "pvp": 89000, "stock": 100,
        "slug": "kit-facial-vitamina-c-bioaqua-5pasos", "nota": "Limpiador, tónico, suero, crema y contorno en una rutina antioxidante."
    },
    {
        "id": "057", "cat": "Cuidado Facial", "sub": "Kits y Rutinas Completas",
        "name": "Kit Facial Purificante y Antiacné con Ácido Salicílico (5 Pasos)",
        "supplier": "23739", "dropiId": "1174240", "costo": 26000, "pvp": 79000, "stock": 169,
        "slug": "kit-facial-antiacne-acido-salicilico-5pasos", "nota": "Tratamiento intensivo de 5 pasos para control de poros y espinillas."
    },
    {
        "id": "058", "cat": "Cuidado Facial", "sub": "Kits y Rutinas Completas",
        "name": "Kit Regenerador Antienvejecimiento Retinol Bioaqua (4 Pasos)",
        "supplier": "98", "dropiId": "2230819", "costo": 38000, "pvp": 99000, "stock": 195,
        "slug": "kit-facial-retinol-antienvejecimiento-4pasos", "nota": "Rutina nocturna regeneradora antiedad con 4 productos activos."
    },
    {
        "id": "059", "cat": "Cuidado Facial", "sub": "Kits y Rutinas Completas",
        "name": "Kit Restaurador de Barrera Cutánea Centella Asiática (4 Pasos)",
        "supplier": "98", "dropiId": "2232979", "costo": 38000, "pvp": 99000, "stock": 187,
        "slug": "kit-restaurador-centella-asiatica-4pasos", "nota": "Protocolo reparador calmante para pieles sensibles o con rojeces."
    },
    {
        "id": "060", "cat": "Cuidado Facial", "sub": "Exfoliantes y Peeling Químico",
        "name": "Exfoliante Dermo-Pulidor Facial con Extracto Botánico 190ml",
        "supplier": "231816", "dropiId": "970396", "costo": 35000, "pvp": 79000, "stock": 100,
        "slug": "exfoliante-dermo-pulidor-facial-botanico-190ml", "nota": "Microcristales minerales biocompatibles que pulen el relieve epidérmico."
    }
]

def main():
    print(f"Cargando libro: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb['Selección_Productos']

    # Fila 32 a 61 (2 + 30)
    for idx, p in enumerate(FACIAL_PRODUCTS, start=32):
        ws.cell(row=idx, column=1).value = p["id"]
        ws.cell(row=idx, column=2).value = "✅ Aprobado"
        ws.cell(row=idx, column=3).value = p["cat"]
        ws.cell(row=idx, column=4).value = p["sub"]
        ws.cell(row=idx, column=5).value = p["name"]
        ws.cell(row=idx, column=6).value = f"Proveedor Verificado #{p['supplier']}"
        ws.cell(row=idx, column=7).value = p["dropiId"]
        ws.cell(row=idx, column=8).value = f"https://app.dropi.co/products/{p['dropiId']}"
        ws.cell(row=idx, column=9).value = p["costo"]
        ws.cell(row=idx, column=10).value = p["pvp"]
        ws.cell(row=idx, column=11).value = f"=(J{idx}-I{idx})/J{idx}"
        ws.cell(row=idx, column=12).value = f"=J{idx}-I{idx}"
        ws.cell(row=idx, column=13).value = "✅ Sí"
        ws.cell(row=idx, column=14).value = 5
        ws.cell(row=idx, column=15).value = "Alta"
        ws.cell(row=idx, column=16).value = "Sí"
        ws.cell(row=idx, column=17).value = p["stock"]
        ws.cell(row=idx, column=18).value = 2
        ws.cell(row=idx, column=19).value = 250
        ws.cell(row=idx, column=20).value = "Media"
        ws.cell(row=idx, column=21).value = "📈 Subiendo"
        ws.cell(row=idx, column=22).value = 9.0
        ws.cell(row=idx, column=23).value = p["nota"]
        ws.cell(row=idx, column=24).value = "2026-10-10"
        ws.cell(row=idx, column=25).value = "Sí"
        ws.cell(row=idx, column=26).value = "2026-10-10"
        ws.cell(row=idx, column=27).value = f"https://mesolabpro.com.co/producto/{p['slug']}"

    wb.save(EXCEL_PATH)
    print("✓ ¡30 productos de Cuidado Facial registrados con éxito en Google Drive!")

if __name__ == "__main__":
    main()
