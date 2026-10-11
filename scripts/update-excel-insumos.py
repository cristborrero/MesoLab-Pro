#!/usr/bin/env python3
"""
Actualiza las filas 152 a 157 de Mesolabpro_Plantilla_Seleccion_Productos.xlsx
con los 6 productos reales y verificados de Insumos Médicos & Desechables.
"""

import os
import openpyxl

EXCEL_GDRIVE = "/Users/cristianborrero/Library/CloudStorage/GoogleDrive-infomesolabpro@gmail.com/My Drive/Mesolabpro/06_Catalogo_y_Productos/02_Estudios_de_Mercado_y_Ganadores/Mesolabpro_Plantilla_Seleccion_Productos.xlsx"
EXCEL_REPO = "docs/Mesolabpro/06_Catalogo_y_Productos/02_Estudios_de_Mercado_y_Ganadores/Mesolabpro_Plantilla_Seleccion_Productos.xlsx"

INSUMOS_PRODUCTS = [
    {
        "id": "151",
        "cat": "Insumos Médicos & Desechables",
        "sub": "Bioseguridad & Guantes",
        "name": "Mascarillas KN95 de Alta Protección x 30 Uds",
        "supplier": "Compra Facil Colombia",
        "dropiId": "2119187",
        "costo": 52500,
        "pvp": 99900,
        "stock": 1200,
        "slug": "mascarillas-kn95-alta-proteccion-30-uds",
        "nota": "Empaque individual hermético. Filtración multicapa BFE ≥95% para cabina y uso clínico.",
    },
    {
        "id": "152",
        "cat": "Insumos Médicos & Desechables",
        "sub": "Bioseguridad & Guantes",
        "name": "Tapabocas Quirúrgico Termosellado Empaque Individual x 50 Uds",
        "supplier": "465911",
        "dropiId": "1696239",
        "costo": 8500,
        "pvp": 39900,
        "stock": 2000,
        "slug": "tapabocas-quirurgico-empaque-individual-50-uds",
        "nota": "Triple capa con filtro meltblown. Empaque individual higiénico en caja por 50 unidades.",
    },
    {
        "id": "153",
        "cat": "Insumos Médicos & Desechables",
        "sub": "Bioseguridad & Guantes",
        "name": "Tapabocas KF94 Coreano de Alta Filtración Estructural",
        "supplier": "55334",
        "dropiId": "286545",
        "costo": 1000,
        "pvp": 24900,
        "stock": 15000,
        "slug": "tapabocas-kf94-coreano-alta-filtracion",
        "nota": "Estructura 3D tipo pez de 4 capas. Mayor volumen respiratorio y antiempañamiento.",
    },
    {
        "id": "154",
        "cat": "Insumos Médicos & Desechables",
        "sub": "Bioseguridad & Guantes",
        "name": "Tapabocas 3D Anatómico Antiempañamiento para Profesionales",
        "supplier": "55334",
        "dropiId": "286537",
        "costo": 3000,
        "pvp": 29900,
        "stock": 28000,
        "slug": "tapabocas-3d-anatomico-antiempanamiento",
        "nota": "Sellado perimetral anatómico. Elimina el empañamiento de lentes y lupas en cabina.",
    },
    {
        "id": "155",
        "cat": "Insumos Médicos & Desechables",
        "sub": "Bioseguridad & Guantes",
        "name": "Tapabocas R5-95 Pediátrico Kids x 10 Unidades",
        "supplier": "55334",
        "dropiId": "286434",
        "costo": 10000,
        "pvp": 24900,
        "stock": 500,
        "slug": "tapabocas-r5-95-kids-10-unidades",
        "nota": "Paquete x 10 uds. Dimensiones reducidas para sellado óptimo en rostros delgados o infantiles.",
    },
    {
        "id": "156",
        "cat": "Insumos Médicos & Desechables",
        "sub": "Desechables de Cabina",
        "name": "Termómetro Infrarrojo Digital Clínico sin Contacto",
        "supplier": "55334",
        "dropiId": "286393",
        "costo": 12000,
        "pvp": 44900,
        "stock": 1899,
        "slug": "termometro-infrarrojo-digital-clinico",
        "nota": "Lectura instantánea en 1 segundo a 3-5 cm. Pantalla LCD con alarma de temperatura.",
    },
]

def update_excel(path):
    if not os.path.exists(path):
        print(f"Archivo no encontrado: {path}")
        return

    print(f"Actualizando: {path}")
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    start_row = 152
    for i, p in enumerate(INSUMOS_PRODUCTS):
        row = start_row + i

        ws[f"A{row}"] = p["id"]
        ws[f"B{row}"] = "Aprobado"
        ws[f"C{row}"] = p["cat"]
        ws[f"D{row}"] = p["sub"]
        ws[f"E{row}"] = p["name"]
        ws[f"F{row}"] = p["supplier"]
        ws[f"G{row}"] = p["dropiId"]
        ws[f"H{row}"] = f"https://app.dropi.co/products/{p['dropiId']}"
        ws[f"I{row}"] = p["costo"]
        ws[f"J{row}"] = p["pvp"]

        # Margen %: =(J - I) / J
        ws[f"K{row}"] = f"=(J{row}-I{row})/J{row}"
        # Margen COP: =J - I
        ws[f"L{row}"] = f"=J{row}-I{row}"

        ws[f"M{row}"] = "Sí"
        ws[f"N{row}"] = 2
        ws[f"O{row}"] = "Alta / CDN"
        ws[f"P{row}"] = "Completa"
        ws[f"Q{row}"] = p["stock"]
        ws[f"R{row}"] = 3  # Días de entrega promedio
        ws[f"S{row}"] = 250  # Peso promedio gramos
        ws[f"T{row}"] = f"https://mesolabpro.com.co/producto/{p['slug']}"
        ws[f"U{row}"] = p["nota"]

        # Formatos numéricos
        ws[f"I{row}"].number_format = "$#,##0"
        ws[f"J{row}"].number_format = "$#,##0"
        ws[f"K{row}"].number_format = "0.0%"
        ws[f"L{row}"].number_format = "$#,##0"

    wb.save(path)
    print(f"✓ Guardado con éxito en {path}")

if __name__ == "__main__":
    update_excel(EXCEL_REPO)
    update_excel(EXCEL_GDRIVE)
