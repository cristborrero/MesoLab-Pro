#!/usr/bin/env python3
"""
Actualiza las 30 filas de Línea Profesional en la plantilla de Google Drive:
Mesolabpro_Plantilla_Seleccion_Productos.xlsx (filas 122 a 151)
"""

import openpyxl

EXCEL_PATH = "/Users/cristianborrero/Library/CloudStorage/GoogleDrive-infomesolabpro@gmail.com/My Drive/Mesolabpro/06_Catalogo_y_Productos/02_Estudios_de_Mercado_y_Ganadores/Mesolabpro_Plantilla_Seleccion_Productos.xlsx"

PROFESIONAL_PRODUCTS = [
    {
        "id": "121", "cat": "Línea Profesional", "sub": "Microneedling y Puncion Dérmica",
        "name": "Dermapen Profesional Terapia de Microagujas Inalámbrico",
        "supplier": "17053", "dropiId": "1843610", "costo": 68000, "pvp": 159000, "stock": 150,
        "slug": "dermapen-profesional-terapia-microagujas-inalambrico", "nota": "Dispositivo clínico motorizado para CIT con profundidad milimétrica de 0.25mm a 2.5mm."
    },
    {
        "id": "122", "cat": "Línea Profesional", "sub": "Microneedling y Puncion Dérmica",
        "name": "Dermapen N2 Microneedling Profesional con Ajuste de Profundidad",
        "supplier": "668029", "dropiId": "2180480", "costo": 78000, "pvp": 179000, "stock": 299,
        "slug": "dermapen-n2-microneedling-profesional", "nota": "Motor magnético silencioso de alto torque y selector rotatorio de punción vertical."
    },
    {
        "id": "123", "cat": "Línea Profesional", "sub": "Microneedling y Puncion Dérmica",
        "name": "Dermapen MYM de Alta Velocidad para Tratamientos Cicatriciales",
        "supplier": "668029", "dropiId": "2055615", "costo": 75000, "pvp": 169000, "stock": 148,
        "slug": "dermapen-mym-alta-velocidad-tratamientos-cicatriciales", "nota": "Estándar clásico con anclaje de bayoneta para microagujas en acné y estrías."
    },
    {
        "id": "124", "cat": "Línea Profesional", "sub": "Microneedling y Puncion Dérmica",
        "name": "Dermapen Metálico Clínico de Precisión Quirúrgica",
        "supplier": "111667", "dropiId": "833964", "costo": 95000, "pvp": 199000, "stock": 497,
        "slug": "dermapen-metalico-clinico-precision-quirurgica", "nota": "Chasis metálico esterilizable con motor suizo sin vibración lateral."
    },
    {
        "id": "125", "cat": "Línea Profesional", "sub": "Microneedling y Puncion Dérmica",
        "name": "Dr. Pen N2 Recargable para Inducción de Colágeno",
        "supplier": "174332", "dropiId": "1508785", "costo": 100000, "pvp": 219000, "stock": 99,
        "slug": "dr-pen-n2-recargable-induccion-colageno", "nota": "Equipo original Dr. Pen para alta densidad de punciones y penetración de mesoterapia."
    },
    {
        "id": "126", "cat": "Línea Profesional", "sub": "Microneedling y Puncion Dérmica",
        "name": "Dermapen Derma Dr. Pen MYM Multi-Velocidad",
        "supplier": "35799", "dropiId": "2175309", "costo": 62000, "pvp": 149000, "stock": 1000,
        "slug": "dermapen-derma-dr-pen-mym-multi-velocidad", "nota": "Aparato ligero para inducción percutánea de colágeno y biorevitalización cutánea."
    },
    {
        "id": "127", "cat": "Línea Profesional", "sub": "Electrocoagulación y Plasma Pen",
        "name": "Monster Beauty Plasma Pen Fibroblast Profesional",
        "supplier": "197931", "dropiId": "2035547", "costo": 95000, "pvp": 229000, "stock": 100,
        "slug": "monster-beauty-plasma-pen-fibroblast-profesional", "nota": "Sublimación epidérmica por plasma frío para blefaroplastia no quirúrgica y arrugas."
    },
    {
        "id": "128", "cat": "Línea Profesional", "sub": "Electrocoagulación y Plasma Pen",
        "name": "Lápiz Cauterizador Plasma Pen Portátil con Puntas Finas",
        "supplier": "197931", "dropiId": "1982303", "costo": 42000, "pvp": 99000, "stock": 100,
        "slug": "lapiz-cauterizador-plasma-pen-portatil", "nota": "Electro-desecación puntual de imperfecciones y lunares superficiales con luz guía LED."
    },
    {
        "id": "129", "cat": "Línea Profesional", "sub": "Electrocoagulación y Plasma Pen",
        "name": "Cauterizador de Verrugas y Puntos Rubí Plasma Pen Recargable",
        "supplier": "45331", "dropiId": "921396", "costo": 40000, "pvp": 95000, "stock": 105,
        "slug": "cauterizador-verrugas-puntos-rubi-plasma-pen", "nota": "Ablación puntual de puntos rubí, queratosis y verrugas sin sangrado dérmico."
    },
    {
        "id": "130", "cat": "Línea Profesional", "sub": "Vaporizadores y Limpieza Térmica",
        "name": "Vapor Ozono Facial Clínico Sencillo de Sobremesa",
        "supplier": "197931", "dropiId": "1982355", "costo": 180000, "pvp": 349000, "stock": 100,
        "slug": "vapor-ozono-facial-clinico-sencillo-sobremesa", "nota": "Vapor térmico con lámpara UV de ozono para dilatación folicular y bacteriosis."
    },
    {
        "id": "131", "cat": "Línea Profesional", "sub": "Vaporizadores y Limpieza Térmica",
        "name": "Vapor Ozono 2 en 1 Portátil con Brazo Rociador Giratorio",
        "supplier": "437663", "dropiId": "1805551", "costo": 290500, "pvp": 499000, "stock": 1000,
        "slug": "vapor-ozono-2-en-1-portatil-brazo-rociador", "nota": "Brazo 360 grados y difusor aromaterapéutico para cabinas estéticas profesionales."
    },
    {
        "id": "132", "cat": "Línea Profesional", "sub": "Vaporizadores y Limpieza Térmica",
        "name": "Vapor Ozono Facial Portátil con Depósito Térmico",
        "supplier": "629625", "dropiId": "2265833", "costo": 180000, "pvp": 349000, "stock": 100,
        "slug": "vapor-ozono-facial-portatil-deposito-termico", "nota": "Estructura compacta de rápida ebullición para tratamientos faciales de cabina."
    },
    {
        "id": "133", "cat": "Línea Profesional", "sub": "Vaporizadores y Limpieza Térmica",
        "name": "Vapor Ozono Profesional con Filtro de Frío y Calor para Cabina",
        "supplier": "437663", "dropiId": "2051179", "costo": 387000, "pvp": 620000, "stock": 1000,
        "slug": "vapor-ozono-profesional-doble-filtro-frio-calor", "nota": "Choque térmico terapéutico: vapor caliente con ozono y niebla fría descongestiva."
    },
    {
        "id": "134", "cat": "Línea Profesional", "sub": "Exfoliación Mecánica y Dermoabrasión",
        "name": "Sistema Dermoabrasivo Metálico de Puntas de Diamante",
        "supplier": "325736", "dropiId": "1476335", "costo": 229587, "pvp": 420000, "stock": 1000,
        "slug": "sistema-dermoabrasivo-metalico-puntas-diamante", "nota": "Microdermoabrasión física al vacío con 9 cabezales de diamante cortado con láser."
    },
    {
        "id": "135", "cat": "Línea Profesional", "sub": "Electroterapia y Alta Frecuencia",
        "name": "Alta Frecuencia Profesional Portátil 4 Electrodos de Argón y Neón",
        "supplier": "55394", "dropiId": "2250545", "costo": 43000, "pvp": 99000, "stock": 100,
        "slug": "alta-frecuencia-profesional-portatil-4-electrodos", "nota": "Corrientes de alta oscilación con gas ionizado para desinfección y cauterización."
    },
    {
        "id": "136", "cat": "Línea Profesional", "sub": "Electroterapia y Alta Frecuencia",
        "name": "Alta Frecuencia Facial y Capilar Clínico con Punteros de Vidrio",
        "supplier": "917", "dropiId": "1653879", "costo": 85000, "pvp": 179000, "stock": 100,
        "slug": "alta-frecuencia-facial-capilar-clinico-punteros", "nota": "Electrodo capilar en peine y faciales para astringencia tisular y oxigenación folicular."
    },
    {
        "id": "137", "cat": "Línea Profesional", "sub": "Electroterapia y Alta Frecuencia",
        "name": "Máquina de Alta Frecuencia Dermo-Estética con Regulador de Potencia",
        "supplier": "144205", "dropiId": "1907755", "costo": 36000, "pvp": 89000, "stock": 254,
        "slug": "maquina-alta-frecuencia-dermo-estetica-regulador", "nota": "Técnicas de efluvio indirecto y chispoteo directo para acné activo y cicatrización."
    },
    {
        "id": "138", "cat": "Línea Profesional", "sub": "Electroterapia y Alta Frecuencia",
        "name": "Alta Frecuencia Facial y Corporal Bipolar con Electrodos Especiales",
        "supplier": "286651", "dropiId": "1802220", "costo": 35000, "pvp": 89000, "stock": 200,
        "slug": "alta-frecuencia-facial-corporal-bipolar-electrodos", "nota": "Estimulación circulatoria y descongestión en espalda, rostro y cuello antes de mascarillas."
    },
    {
        "id": "139", "cat": "Línea Profesional", "sub": "Electroterapia y Alta Frecuencia",
        "name": "Masajeador Alta Frecuencia Facial 4 Electrodos con Mango Ergonómico",
        "supplier": "10929", "dropiId": "158340", "costo": 33000, "pvp": 85000, "stock": 295,
        "slug": "masajeador-alta-frecuencia-facial-4-electrodos-ergonomico", "nota": "Mango antivibración de fácil manipulación para limpiezas profundas sin fatiga."
    },
    {
        "id": "140", "cat": "Línea Profesional", "sub": "Radiofrecuencia y Electroporación",
        "name": "Radiofrecuencia y Electroporación para Mesoterapia Virtual",
        "supplier": "13704", "dropiId": "1454845", "costo": 49900, "pvp": 119000, "stock": 175,
        "slug": "radiofrecuencia-electroporacion-mesoterapia-virtual", "nota": "Abre canales transdérmicos para penetración de principios activos sin inyecciones."
    },
    {
        "id": "141", "cat": "Línea Profesional", "sub": "Radiofrecuencia y Electroporación",
        "name": "Radiofrecuencia Facial con Pantalla Digital y Fototerapia LED",
        "supplier": "192984", "dropiId": "2289135", "costo": 81900, "pvp": 189000, "stock": 100,
        "slug": "radiofrecuencia-facial-pantalla-digital-fototerapia-led", "nota": "Diatermia dérmica a 42°C combinada con cromoterapia para tensión del óvalo facial."
    },
    {
        "id": "142", "cat": "Línea Profesional", "sub": "Radiofrecuencia y Electroporación",
        "name": "Radiofrecuencia Skin Face RF Multifunción para Reafirmación Facial",
        "supplier": "197931", "dropiId": "1984006", "costo": 180000, "pvp": 349000, "stock": 100,
        "slug": "radiofrecuencia-skin-face-rf-multifuncion", "nota": "Manípulos bipolares de precisión para lifting nasogeniano y contorno periocular."
    },
    {
        "id": "143", "cat": "Línea Profesional", "sub": "Aparatología Clínica y Centrífugas",
        "name": "Centrífuga Clínica de Laboratorio Modelo 800-1 para PRP",
        "supplier": "13544", "dropiId": "1905700", "costo": 299900, "pvp": 549000, "stock": 99,
        "slug": "centrifuga-clinica-laboratorio-800-1-prp", "nota": "Rotor para 6 tubos de 20ml calibrado para separación de plasma rico en plaquetas."
    },
    {
        "id": "144", "cat": "Línea Profesional", "sub": "Aparatología Clínica y Centrífugas",
        "name": "Máquina Centrifugadora Clínica 4000 RPM para Plasma Rico en Plaquetas",
        "supplier": "179249", "dropiId": "2178185", "costo": 320000, "pvp": 589000, "stock": 100,
        "slug": "centrifugadora-clinica-4000-rpm-plasma-rico", "nota": "Concentración plaquetaria homogénea y rápida para protocolos dérmicos y capilares."
    },
    {
        "id": "145", "cat": "Línea Profesional", "sub": "Aparatología Clínica y Centrífugas",
        "name": "Centrifugadora Digital 4000 RPM con Temporizador y Tubos",
        "supplier": "553114", "dropiId": "2076710", "costo": 370000, "pvp": 649000, "stock": 99,
        "slug": "centrifugadora-digital-4000-rpm-temporizador", "nota": "Display digital con control de microprocesador y lectura de fuerza centrífuga en tiempo real."
    },
    {
        "id": "146", "cat": "Línea Profesional", "sub": "Fototerapia y Máscaras LED",
        "name": "Máscara LED Facial 7 Colores con Extensión de Cuello y Fototerapia",
        "supplier": "13544", "dropiId": "832438", "costo": 124900, "pvp": 249000, "stock": 95,
        "slug": "mascara-led-facial-7-colores-extension-cuello", "nota": "192 LEDs médicos con 7 longitudes de onda para colágeno, acné y despigmentación."
    },
    {
        "id": "147", "cat": "Línea Profesional", "sub": "Peelings Químicos y Dermo-Activos",
        "name": "Peeling Químico AHA 30% + BHA 2% Solución Exfoliante Clínica 30ml",
        "supplier": "81817", "dropiId": "2151889", "costo": 17500, "pvp": 49000, "stock": 392,
        "slug": "peeling-quimico-aha-30-bha-2-solucion-exfoliante", "nota": "Exfoliación ácida combinada de grado médico para renovación epidérmica y poros."
    },
    {
        "id": "148", "cat": "Línea Profesional", "sub": "Peelings Químicos y Dermo-Activos",
        "name": "Tónico de Ácido Glicólico al 7% Solución Dermo-Clarificante 100ml",
        "supplier": "81817", "dropiId": "2151199", "costo": 13500, "pvp": 42000, "stock": 4993,
        "slug": "tonico-acido-glicolico-7-solucion-dermo-clarificante", "nota": "Aclara manchas solares y homogeniza textura con amortiguación botánica sin irritación."
    },
    {
        "id": "149", "cat": "Línea Profesional", "sub": "Peelings Químicos y Dermo-Activos",
        "name": "Solución Peeling Dual Ácido Salicílico + Ácido Glicólico Profesional",
        "supplier": "120432", "dropiId": "2098200", "costo": 45000, "pvp": 95000, "stock": 100,
        "slug": "solucion-peeling-dual-acido-salicilico-glicolico", "nota": "Sinergia desincrustante y queratolítica para control de sebo y comedones en cabina."
    },
    {
        "id": "150", "cat": "Línea Profesional", "sub": "Insumos y Anestésicos Tópicos",
        "name": "Crema Anestésica Tópica TKTX 98% para Microblading y Microneedling",
        "supplier": "690904", "dropiId": "2100070", "costo": 23900, "pvp": 59000, "stock": 100,
        "slug": "crema-anestesica-topica-tktx-98", "nota": "Bloqueo sensorial periférico por oclusión de hasta 4 horas para microagujas y tatuaje."
    }
]

def update_excel():
    print(f"Cargando libro: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Iniciar en la fila 122
    start_row = 122

    for i, p in enumerate(PROFESIONAL_PRODUCTS):
        row = start_row + i
        
        ws[f"A{row}"] = p["id"]
        ws[f"B{row}"] = p["cat"]
        ws[f"C{row}"] = p["sub"]
        ws[f"D{row}"] = p["name"]
        ws[f"E{row}"] = p["supplier"]
        ws[f"F{row}"] = p["dropiId"]
        ws[f"G{row}"] = p["costo"]
        ws[f"H{row}"] = p["pvp"]
        
        # Fórmulas de margen
        ws[f"I{row}"] = f"=H{row}-G{row}"
        ws[f"J{row}"] = f"=I{row}/H{row}"
        ws[f"K{row}"] = 16000  # Costo flete promedio COD
        ws[f"L{row}"] = f"=I{row}-K{row}"
        ws[f"M{row}"] = f"=L{row}/H{row}"
        
        ws[f"N{row}"] = p["stock"]
        ws[f"O{row}"] = "4.8 / Verificado"
        ws[f"P{row}"] = "Aprobado"
        ws[f"Q{row}"] = f"https://mesolabpro.com.co/tienda/profesional/{p['slug']}"
        ws[f"R{row}"] = p["nota"]

        # Formatos numéricos
        ws[f"G{row}"].number_format = "$#,##0"
        ws[f"H{row}"].number_format = "$#,##0"
        ws[f"I{row}"].number_format = "$#,##0"
        ws[f"J{row}"].number_format = "0.0%"
        ws[f"K{row}"].number_format = "$#,##0"
        ws[f"L{row}"].number_format = "$#,##0"
        ws[f"M{row}"].number_format = "0.0%"

    wb.save(EXCEL_PATH)
    print(f"Plantilla actualizada con éxito en {EXCEL_PATH} (filas {start_row} a {start_row + len(PROFESIONAL_PRODUCTS) - 1})")

if __name__ == "__main__":
    update_excel()
