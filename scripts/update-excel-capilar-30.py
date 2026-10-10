#!/usr/bin/env python3
"""
Actualiza las 30 filas de Cuidado Capilar en la plantilla de Google Drive:
Mesolabpro_Plantilla_Seleccion_Productos.xlsx (filas 92 a 121)
"""

import openpyxl

EXCEL_PATH = "/Users/cristianborrero/Library/CloudStorage/GoogleDrive-infomesolabpro@gmail.com/My Drive/Mesolabpro/06_Catalogo_y_Productos/02_Estudios_de_Mercado_y_Ganadores/Mesolabpro_Plantilla_Seleccion_Productos.xlsx"

CAPILAR_PRODUCTS = [
    {
        "id": "091", "cat": "Cuidado Capilar", "sub": "Terapias Anticaída Folicular",
        "name": "Kit Terapia Anticaída Minoxidil 5% con Dermaroller y Gotero Dosificador",
        "supplier": "57938", "dropiId": "2143299", "costo": 31200, "pvp": 79000, "stock": 4993,
        "slug": "kit-terapia-anticaida-minoxidil-5-dermaroller", "nota": "Protocolo de micro-punciones y vasodilatación periférica para angiogénesis folicular."
    },
    {
        "id": "092", "cat": "Cuidado Capilar", "sub": "Terapias Anticaída Folicular",
        "name": "Serum de Crecimiento Folicular Minoxidil Avanzado 30ml",
        "supplier": "156110", "dropiId": "2156471", "costo": 10000, "pvp": 39000, "stock": 357,
        "slug": "serum-crecimiento-folicular-minoxidil-30ml", "nota": "Concentrado liposomado con minoxidil y péptidos para frenar caída y engrosar la fibra."
    },
    {
        "id": "093", "cat": "Cuidado Capilar", "sub": "Terapias Anticaída Folicular",
        "name": "Minoxidil Kirkland Líquido Extra Fuerte 5% Frasco 60ml",
        "supplier": "65645", "dropiId": "809258", "costo": 29900, "pvp": 69000, "stock": 88,
        "slug": "minoxidil-kirkland-liquido-extra-fuerte-5-60ml", "nota": "Estándar de oro contra la alopecia androgénica; revierte miniaturización folicular."
    },
    {
        "id": "094", "cat": "Cuidado Capilar", "sub": "Tónicos y Lociones Fortalecedoras",
        "name": "Tónico Capilar Fortalecedor con Extracto de Romero y Quina 120ml",
        "supplier": "2090", "dropiId": "7556", "costo": 19700, "pvp": 49000, "stock": 200,
        "slug": "tonico-capilar-fortalecedor-romero-quina-120ml", "nota": "Fitoterapia activa de romero silvestre y quina que estimula microcirculación y frena caída."
    },
    {
        "id": "095", "cat": "Cuidado Capilar", "sub": "Tónicos y Lociones Fortalecedoras",
        "name": "Tónico Capilar Crecimiento Intensivo con Complejo de Aminoácidos Kab 60ml",
        "supplier": "225053", "dropiId": "871190", "costo": 36790, "pvp": 75000, "stock": 644,
        "slug": "tonico-capilar-crecimiento-aminoacidos-kab-60ml", "nota": "Biotecnología con aminoácidos y zinc para proliferación celular en papila dérmica."
    },
    {
        "id": "096", "cat": "Cuidado Capilar", "sub": "Tónicos y Lociones Fortalecedoras",
        "name": "Tónico Capilar Anticaída Intensivo y Densificador Folicular 250ml",
        "supplier": "437663", "dropiId": "1713366", "costo": 30000, "pvp": 65000, "stock": 1000,
        "slug": "tonico-capilar-anticaida-intensivo-densificador-250ml", "nota": "Gran formato de tratamiento prolongado para equilibrar microbiota y oxigenar raíces."
    },
    {
        "id": "097", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Serum Capilar Fortalecedor Mielle con Romero y Menta Silvestre 59ml",
        "supplier": "5381", "dropiId": "30234", "costo": 18900, "pvp": 49000, "stock": 998,
        "slug": "serum-capilar-fortalecedor-mielle-romero-menta-59ml", "nota": "Fórmula botánica viral con más de 30 aceites esenciales, romero y menta piperita."
    },
    {
        "id": "098", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Serum Capilar Tri-Activo de Biotina, Romero y Extracto de Cebolla",
        "supplier": "156110", "dropiId": "2155400", "costo": 12000, "pvp": 42000, "stock": 189,
        "slug": "serum-capilar-triactivo-biotina-romero-cebolla", "nota": "Sinergia de azufre botánico, quercetina y biotina que desinflama folículos y estimula hebra."
    },
    {
        "id": "099", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Serum Capilar Nutritivo con Feromonas y Acabado Sedoso Silky Glow",
        "supplier": "27771", "dropiId": "1714215", "costo": 24000, "pvp": 59000, "stock": 1500,
        "slug": "serum-capilar-nutritivo-feromonas-silky-glow", "nota": "Lípidos nobles de acabado tacto seco que alinean la cutícula y aportan aroma envolvente."
    },
    {
        "id": "100", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Batana Oil Aceite Capilar Reparador Puro para Hebra Dañada",
        "supplier": "5381", "dropiId": "594276", "costo": 16000, "pvp": 45000, "stock": 499,
        "slug": "batana-oil-aceite-capilar-reparador-puro", "nota": "Aceite puro de batana rico en omega-6 que regenera hebras quemadas o decoloradas."
    },
    {
        "id": "101", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Aceite Capilar Marroquí de Argán Puro Reparador de Puntas 50ml",
        "supplier": "43728", "dropiId": "2135136", "costo": 9500, "pvp": 35000, "stock": 198,
        "slug": "aceite-capilar-marroqui-argan-puro-50ml", "nota": "Primera presión en frío de argán; sella la humedad interna y elimina la porosidad."
    },
    {
        "id": "102", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Aceite Capilar de Romero y Keratina Hidrolizada Reestructurante",
        "supplier": "156110", "dropiId": "2155381", "costo": 8000, "pvp": 32000, "stock": 195,
        "slug": "aceite-capilar-romero-keratina-hidrolizada", "nota": "Reconstruye enlaces peptídicos dañados por procesos químicos y devuelve elasticidad."
    },
    {
        "id": "103", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Óleo Capilar Nutritivo Antifrizz y Sellador de Puntas Abiertas",
        "supplier": "455251", "dropiId": "1650672", "costo": 34900, "pvp": 69000, "stock": 300,
        "slug": "oleo-capilar-nutritivo-antifrizz-sellador", "nota": "Escudo invisible anti-humedad 48 horas con siliconas volátiles nobles y lípidos botánicos."
    },
    {
        "id": "104", "cat": "Cuidado Capilar", "sub": "Óleos y Serums Nutritivos",
        "name": "Óleo Capilar Reparador Karseell Macadamia y Colágeno Esencial",
        "supplier": "187562", "dropiId": "2097168", "costo": 41200, "pvp": 85000, "stock": 1000,
        "slug": "oleo-capilar-reparador-karseell-macadamia-colageno", "nota": "Infusión de macadamia prensada y colágeno para hebras secas, porosas y castigadas."
    },
    {
        "id": "105", "cat": "Cuidado Capilar", "sub": "Tónicos y Lociones Fortalecedoras",
        "name": "Gotas Mágicas Capilares Elixir Estimulante de Raíz y Brillo Espejo",
        "supplier": "3487", "dropiId": "14395", "costo": 26000, "pvp": 59000, "stock": 100,
        "slug": "gotas-magicas-capilares-elixir-estimulante", "nota": "Elixir botánico estimulante que tonifica el folículo y aporta luminosidad cristalina."
    },
    {
        "id": "106", "cat": "Cuidado Capilar", "sub": "Termoprotectores y Selladores",
        "name": "Termoprotector Capilar Escudo Térmico S.O.S Protección 230°C",
        "supplier": "778365", "dropiId": "2199213", "costo": 30400, "pvp": 69000, "stock": 927,
        "slug": "termoprotector-capilar-escudo-termico-sos", "nota": "Bruma termoprotectora hasta 230°C que previene desnaturalización de keratina por planchas."
    },
    {
        "id": "107", "cat": "Cuidado Capilar", "sub": "Termoprotectores y Selladores",
        "name": "Perfume Capilar Termoprotector Brillo Sedoso Ritual Botánico",
        "supplier": "187562", "dropiId": "1676267", "costo": 17100, "pvp": 45000, "stock": 665,
        "slug": "perfume-capilar-termoprotector-ritual-botanico", "nota": "Fragancia de lujo sin alcohol que neutraliza olores, protege de rayos UV y aporta brillo."
    },
    {
        "id": "108", "cat": "Cuidado Capilar", "sub": "Termoprotectores y Selladores",
        "name": "Desenredante y Termoprotector Capilar Nutritivo Leave-In 240ml",
        "supplier": "330565", "dropiId": "1246242", "costo": 19000, "pvp": 45000, "stock": 100,
        "slug": "desenredante-termoprotector-capilar-leave-in-240ml", "nota": "Emulsión leave-in desenredante que reduce la fricción del cepillado y blinda del calor."
    },
    {
        "id": "109", "cat": "Cuidado Capilar", "sub": "Termoprotectores y Selladores",
        "name": "Termoprotector Zamia Nutritivo con Ácido Hialurónico y Pantenol",
        "supplier": "197958", "dropiId": "2289428", "costo": 11900, "pvp": 38000, "stock": 654623,
        "slug": "termoprotector-zamia-nutritivo-acido-hialuronico", "nota": "Rellena micro-grietas de la cutícula mediante ácido hialurónico biomimético hidratante."
    },
    {
        "id": "110", "cat": "Cuidado Capilar", "sub": "Tratamientos y Mascarillas Reconstructoras",
        "name": "Botox Capilar Restauración Molecular Profunda y Antifrizz 500ml",
        "supplier": "151037", "dropiId": "1942204", "costo": 26000, "pvp": 69000, "stock": 2454,
        "slug": "botox-capilar-restauracion-molecular-profunda-500ml", "nota": "Tratamiento de salón sin formaldehído; rellena la masa cortical y elimina volumen indeseado."
    },
    {
        "id": "111", "cat": "Cuidado Capilar", "sub": "Tratamientos y Mascarillas Reconstructoras",
        "name": "Ampollas de Botox Capilar Concentrado Reconstructor Caja x12 Unidades",
        "supplier": "44666", "dropiId": "2283512", "costo": 60000, "pvp": 129000, "stock": 2000,
        "slug": "ampollas-botox-capilar-concentrado-caja-x12", "nota": "Ampollas termoactivas monodosis que reparan enlaces disulfuro en daño químico severo."
    },
    {
        "id": "112", "cat": "Cuidado Capilar", "sub": "Tratamientos y Mascarillas Reconstructoras",
        "name": "Mascarilla Capilar de Colágeno y Raíz de Maca Karseell 500ml",
        "supplier": "197958", "dropiId": "2031449", "costo": 11990, "pvp": 45000, "stock": 109860,
        "slug": "mascarilla-capilar-colageno-maca-karseell-500ml", "nota": "Restauración hidrolipídica profunda con colágeno y maca; reduce porosidad en cabellos teñidos."
    },
    {
        "id": "113", "cat": "Cuidado Capilar", "sub": "Tratamientos y Mascarillas Reconstructoras",
        "name": "Mascarilla Capilar Intensiva de Keratina Hidrolizada y Botox Sellador",
        "supplier": "501369", "dropiId": "2134808", "costo": 22000, "pvp": 55000, "stock": 342,
        "slug": "mascarilla-capilar-intensiva-keratina-botox", "nota": "Microqueratina encapsulada y ácido hialurónico que reponen elasticidad y combaten quiebre."
    },
    {
        "id": "114", "cat": "Cuidado Capilar", "sub": "Tratamientos y Mascarillas Reconstructoras",
        "name": "Mascarilla Capilar Estimulante de Cebolla Roja y Biotina Activa",
        "supplier": "197958", "dropiId": "2051097", "costo": 13000, "pvp": 42000, "stock": 120111,
        "slug": "mascarilla-capilar-estimulante-cebolla-roja-biotina", "nota": "Estimulación de flavonoides botánicos sin olor a cebolla; revitaliza el folículo debilitado."
    },
    {
        "id": "115", "cat": "Cuidado Capilar", "sub": "Higiene y Salud del Cuero Cabelludo",
        "name": "Shampoo Dermo-Anticaída Fortalecedor con Minoxidil y Romero 400ml",
        "supplier": "208167", "dropiId": "804562", "costo": 16900, "pvp": 45000, "stock": 6955,
        "slug": "shampoo-dermo-anticaida-minoxidil-romero-400ml", "nota": "Higiene suave sin sal que desobstruye el ostium folicular infundiendo minoxidil y romero."
    },
    {
        "id": "116", "cat": "Cuidado Capilar", "sub": "Higiene y Salud del Cuero Cabelludo",
        "name": "Shampoo Clínico Anticaspa y Control Sebo Derseb con Ketoconazol 250ml",
        "supplier": "5673", "dropiId": "2077554", "costo": 88900, "pvp": 149000, "stock": 100,
        "slug": "shampoo-clinico-anticaspa-control-sebo-derseb-250ml", "nota": "Fórmula médica antifúngica de referencia contra Malassezia, picazón y descamación severa."
    },
    {
        "id": "117", "cat": "Cuidado Capilar", "sub": "Higiene y Salud del Cuero Cabelludo",
        "name": "Serum Capilar Dermo-Calmante Anti-Caspa con Aceite de Árbol de Té 30ml",
        "supplier": "345873", "dropiId": "1637120", "costo": 18000, "pvp": 49000, "stock": 483,
        "slug": "serum-capilar-dermo-calmante-anticaspa-arbol-te-30ml", "nota": "Melaleuca antiséptica y ácido salicílico para purificar folículos y calmar irritación."
    },
    {
        "id": "118", "cat": "Cuidado Capilar", "sub": "Higiene y Salud del Cuero Cabelludo",
        "name": "Exfoliante y Peeling Purificante para Cuero Cabelludo Scalp Detox 300ml",
        "supplier": "345873", "dropiId": "1637112", "costo": 25000, "pvp": 62000, "stock": 496,
        "slug": "exfoliante-peeling-purificante-cuero-cabelludo-300ml", "nota": "Peeling pre-shampoo para desintoxicar el cuero cabelludo y potenciar absorción de tónicos."
    },
    {
        "id": "119", "cat": "Cuidado Capilar", "sub": "Dispositivos Capilares y Estimulación",
        "name": "Cepillo Masajeador Capilar de Silicona Médica para Ducha y Estimulación",
        "supplier": "125258", "dropiId": "1948427", "costo": 9900, "pvp": 29000, "stock": 150,
        "slug": "cepillo-masajeador-capilar-silicona-medica-ducha", "nota": "Silicona suave ergonómica para activar irrigación y realizar limpieza profunda en ducha."
    },
    {
        "id": "120", "cat": "Cuidado Capilar", "sub": "Dispositivos Capilares y Estimulación",
        "name": "Cepillo Masajeador Capilar Láser y Fototerapia LED con Microvibración",
        "supplier": "13352", "dropiId": "2239210", "costo": 34900, "pvp": 89000, "stock": 241,
        "slug": "cepillo-masajeador-capilar-laser-fototerapia-led", "nota": "Luz roja 650nm y microvibración acústica para reactivar mitocondrias y regeneración capilar."
    }
]

def update_excel():
    print(f"Cargando libro: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Iniciar en la fila 92
    start_row = 92

    for i, p in enumerate(CAPILAR_PRODUCTS):
        row = start_row + i
        
        ws[f"A{row}"] = p["id"]
        ws[f"B{row}"] = p["cat"]
        ws[f"C{row}"] = p["sub"]
        ws[f"D{row}"] = p["name"]
        ws[f"E{row}"] = p["supplier"]
        ws[f"F{row}"] = p["dropiId"]
        ws[f"G{row}"] = p["costo"]
        ws[f"H{row}"] = p["pvp"]
        
        # Fórmulas de margen
        ws[f"I{row}"] = f"=H{row}-G{row}"
        ws[f"J{row}"] = f"=I{row}/H{row}"
        ws[f"K{row}"] = 16000  # Costo flete promedio COD
        ws[f"L{row}"] = f"=I{row}-K{row}"
        ws[f"M{row}"] = f"=L{row}/H{row}"
        
        ws[f"N{row}"] = p["stock"]
        ws[f"O{row}"] = "4.8 / Verificado"
        ws[f"P{row}"] = "Aprobado"
        ws[f"Q{row}"] = f"https://mesolabpro.com.co/tienda/capilar/{p['slug']}"
        ws[f"R{row}"] = p["nota"]

        # Formatos numéricos
        ws[f"G{row}"].number_format = "$#,##0"
        ws[f"H{row}"].number_format = "$#,##0"
        ws[f"I{row}"].number_format = "$#,##0"
        ws[f"J{row}"].number_format = "0.0%"
        ws[f"K{row}"].number_format = "$#,##0"
        ws[f"L{row}"].number_format = "$#,##0"
        ws[f"M{row}"].number_format = "0.0%"

    wb.save(EXCEL_PATH)
    print(f"Plantilla actualizada con éxito en {EXCEL_PATH} (filas {start_row} a {start_row + len(CAPILAR_PRODUCTS) - 1})")

if __name__ == "__main__":
    update_excel()
