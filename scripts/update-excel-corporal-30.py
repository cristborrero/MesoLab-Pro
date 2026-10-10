#!/usr/bin/env python3
"""
Actualiza las 30 filas de Cuidado Corporal en la plantilla de Google Drive:
Mesolabpro_Plantilla_Seleccion_Productos.xlsx (filas 62 a 91)
"""

import openpyxl

EXCEL_PATH = "/Users/cristianborrero/Library/CloudStorage/GoogleDrive-infomesolabpro@gmail.com/My Drive/Mesolabpro/06_Catalogo_y_Productos/02_Estudios_de_Mercado_y_Ganadores/Mesolabpro_Plantilla_Seleccion_Productos.xlsx"

CORPORAL_PRODUCTS = [
    {
        "id": "061", "cat": "Cuidado Corporal", "sub": "Reductores y Termoactivos",
        "name": "Gel Reductor Abdominal Termoactivo con Cafeína Sculpt 100ml",
        "supplier": "29636", "dropiId": "2094233", "costo": 10000, "pvp": 39000, "stock": 5070,
        "slug": "gel-reductor-abdominal-termoactivo-sculpt-100ml", "nota": "Acción térmica localizada que estimula la microcirculación y degradación lipídica."
    },
    {
        "id": "062", "cat": "Cuidado Corporal", "sub": "Reductores y Termoactivos",
        "name": "Gel Reductor y Moldeador Corporal Profesional 500ml",
        "supplier": "345873", "dropiId": "1637130", "costo": 29000, "pvp": 69000, "stock": 492,
        "slug": "gel-reductor-moldeador-corporal-profesional-500ml", "nota": "Formato profesional para masajes reductores intensivos de cintura y muslos."
    },
    {
        "id": "063", "cat": "Cuidado Corporal", "sub": "Reductores y Termoactivos",
        "name": "Gel Caliente Termogénico Lipolítico MW 250ml",
        "supplier": "10135", "dropiId": "71361", "costo": 29000, "pvp": 65000, "stock": 332,
        "slug": "gel-caliente-termogenico-lipolitico-mw-250ml", "nota": "Efecto vasodilatador profundo con extractos vegetales que eleva temperatura tisular."
    },
    {
        "id": "064", "cat": "Cuidado Corporal", "sub": "Reductores y Termoactivos",
        "name": "Gel Caliente Reductor Moldeador Plus con Centella",
        "supplier": "19662", "dropiId": "1617898", "costo": 9800, "pvp": 38000, "stock": 1716,
        "slug": "gel-caliente-reductor-moldeador-plus-centella", "nota": "Estimulante térmico ligero enriquecido con centella asiática para drenaje adiposo."
    },
    {
        "id": "065", "cat": "Cuidado Corporal", "sub": "Reafirmantes y Cuello/Escote",
        "name": "Crema Reafirmante Tensora con Manteca de Karité Palmer's",
        "supplier": "156110", "dropiId": "2174419", "costo": 13000, "pvp": 45000, "stock": 200,
        "slug": "crema-reafirmante-tensora-manteca-karite-palmers", "nota": "Manteca de karité, colágeno y elastina pura para tonificar tejidos post-parto o dieta."
    },
    {
        "id": "066", "cat": "Cuidado Corporal", "sub": "Reafirmantes y Cuello/Escote",
        "name": "Crema Tensora Efecto Lifting para Cuello y Escote 100g",
        "supplier": "205895", "dropiId": "1647396", "costo": 9000, "pvp": 39000, "stock": 22820,
        "slug": "crema-tensora-efecto-lifting-cuello-escote-100g", "nota": "Péptidos biomiméticos que combaten la flacidez, descolgamiento y anillos de Venus."
    },
    {
        "id": "067", "cat": "Cuidado Corporal", "sub": "Reafirmantes y Cuello/Escote",
        "name": "Crema Reafirmante Corporal y Glúteos Firmsta 200g",
        "supplier": "54262", "dropiId": "1210719", "costo": 18900, "pvp": 55000, "stock": 7670,
        "slug": "crema-reafirmante-corporal-gluteos-firmsta-200g", "nota": "Activos botánicos tensores y estimuladores de colágeno para glúteos y muslos."
    },
    {
        "id": "068", "cat": "Cuidado Corporal", "sub": "Reafirmantes y Cuello/Escote",
        "name": "Crema Corporal Reafirmante Hidratación Profunda 500ml",
        "supplier": "197958", "dropiId": "2287234", "costo": 11000, "pvp": 42000, "stock": 526598,
        "slug": "crema-corporal-reafirmante-hidratacion-profunda-500ml", "nota": "Emulsión corporal de gran formato enriquecida con aminoácidos y ceramidas."
    },
    {
        "id": "069", "cat": "Cuidado Corporal", "sub": "Anticelulitis y Drenaje",
        "name": "Crema Anticelulítica Adeus con Extracto de Café y Ginkgo",
        "supplier": "5381", "dropiId": "280684", "costo": 11900, "pvp": 45000, "stock": 562,
        "slug": "crema-anticelulitica-adeus-cafe-ginkgo", "nota": "Reduce apariencia de piel de naranja mejorando textura cutánea y retención hídrica."
    },
    {
        "id": "070", "cat": "Cuidado Corporal", "sub": "Anticelulitis y Drenaje",
        "name": "Aceite Reafirmante Corporal Anticelulítico con Vitamina E 100ml",
        "supplier": "29636", "dropiId": "2118012", "costo": 14000, "pvp": 49000, "stock": 1104,
        "slug": "aceite-reafirmante-corporal-anticelulitico-vitamina-e-100ml", "nota": "Aceite seco de rápida absorción para drenaje linfático manual y activación circulatoria."
    },
    {
        "id": "071", "cat": "Cuidado Corporal", "sub": "Anticelulitis y Drenaje",
        "name": "Rodillo Masajeador 360 Drenante y Anticelulítico con 9 Ruedas",
        "supplier": "6006", "dropiId": "2124344", "costo": 18900, "pvp": 49000, "stock": 112,
        "slug": "rodillo-masajeador-360-drenante-anticelulitico", "nota": "Dispositivo envolvente de maderoterapia moderna para romper nódulos y drenar linfa."
    },
    {
        "id": "072", "cat": "Cuidado Corporal", "sub": "Anticelulitis y Drenaje",
        "name": "Masajeador Corporal Anticelulítico Eléctrico 2.0 con Infrarrojos",
        "supplier": "64834", "dropiId": "337126", "costo": 46900, "pvp": 99000, "stock": 156,
        "slug": "masajeador-corporal-anticelulitico-electrico-infrarrojo-2", "nota": "Cabezales rotativos con calor infrarrojo para alisar celulitis y tonificar tejido."
    },
    {
        "id": "073", "cat": "Cuidado Corporal", "sub": "Antiestrías y Cicatrices",
        "name": "Loción de Masaje para Estrías con Manteca de Cacao Palmer's 250ml",
        "supplier": "796149", "dropiId": "2173089", "costo": 37900, "pvp": 79000, "stock": 980,
        "slug": "locion-masaje-estrias-manteca-cacao-palmers-250ml", "nota": "Fórmula con manteca de cacao pura, colágeno y elastina para estrías de embarazo."
    },
    {
        "id": "074", "cat": "Cuidado Corporal", "sub": "Antiestrías y Cicatrices",
        "name": "Crema Antiestrías y Cicatrices SkinRepair Regenerativa",
        "supplier": "197958", "dropiId": "2120358", "costo": 8000, "pvp": 35000, "stock": 587395,
        "slug": "crema-antiestrias-cicatrices-skinrepair-regenerativa", "nota": "Tratamiento emoliente que acelera la renovación tisular y suaviza estrías."
    },
    {
        "id": "075", "cat": "Cuidado Corporal", "sub": "Antiestrías y Cicatrices",
        "name": "Rutina Corporal Antiestrías y Firmeza Dúo Reparador",
        "supplier": "434826", "dropiId": "2190247", "costo": 23000, "pvp": 69000, "stock": 1000,
        "slug": "rutina-corporal-antiestrias-firmeza-duo-reparador", "nota": "Tratamiento dual para zonas propensas a ruptura dérmica (vientre, senos, caderas)."
    },
    {
        "id": "076", "cat": "Cuidado Corporal", "sub": "Exfoliantes y Pulidores",
        "name": "Exfoliante Corporal Dermo-Pulidor con Microgránulos 200ml",
        "supplier": "345873", "dropiId": "1637128", "costo": 28000, "pvp": 69000, "stock": 500,
        "slug": "exfoliante-corporal-dermopulidor-microgranulos-200ml", "nota": "Gránulos biodegradables de pulido que afinan el poro y preparan la piel para activos."
    },
    {
        "id": "077", "cat": "Cuidado Corporal", "sub": "Exfoliantes y Pulidores",
        "name": "Body Scrub Exfoliante Corporal Botánico Dermanat",
        "supplier": "323312", "dropiId": "1386131", "costo": 22647, "pvp": 59000, "stock": 100,
        "slug": "body-scrub-exfoliante-corporal-botanico-dermanat", "nota": "Scrub con extractos botánicos y cristales minerales para desintoxicar y suavizar."
    },
    {
        "id": "078", "cat": "Cuidado Corporal", "sub": "Exfoliantes y Pulidores",
        "name": "Exfoliante Corporal Sedoso Truly Rainbow Glow 200ml",
        "supplier": "525739", "dropiId": "1898089", "costo": 18000, "pvp": 55000, "stock": 1904,
        "slug": "exfoliante-corporal-sedoso-truly-rainbow-glow-200ml", "nota": "Mousse exfoliante batida con colágeno vegetal y azúcar refinado no agresivo."
    },
    {
        "id": "079", "cat": "Cuidado Corporal", "sub": "Exfoliantes y Pulidores",
        "name": "Jabón Dermo-Exfoliante Masajeador con Micropartículas 120g",
        "supplier": "43339", "dropiId": "1812615", "costo": 9970, "pvp": 29000, "stock": 2094,
        "slug": "jabon-dermoexfoliante-masajeador-microparticulas-120g", "nota": "Barra de masaje con relieves esféricos que exfolia y activa la microcirculación."
    },
    {
        "id": "080", "cat": "Cuidado Corporal", "sub": "Aclarantes y Zonas Íntimas",
        "name": "Serum Aclarante y Despigmentante Corporal Bioaqua 30ml",
        "supplier": "157185", "dropiId": "781840", "costo": 15500, "pvp": 45000, "stock": 261,
        "slug": "serum-aclarante-despigmentante-corporal-bioaqua-30ml", "nota": "Arbutina y niacinamida para aclarar manchas en axilas, entrepierna, codos y cuello."
    },
    {
        "id": "081", "cat": "Cuidado Corporal", "sub": "Aclarantes y Zonas Íntimas",
        "name": "Gel Despigmentante Íntimo de Alta Tolerancia Piel de Oro 60ml",
        "supplier": "507182", "dropiId": "1750479", "costo": 30900, "pvp": 69000, "stock": 247,
        "slug": "gel-despigmentante-intimo-alta-tolerancia-piel-de-oro-60ml", "nota": "Fórmula a pH fisiológico para zonas delicadas libre de hidroquinona o alcohol."
    },
    {
        "id": "082", "cat": "Cuidado Corporal", "sub": "Aclarantes y Zonas Íntimas",
        "name": "Serum Despigmentante Corporal Smooth Legend 90ml",
        "supplier": "525739", "dropiId": "1942028", "costo": 24500, "pvp": 59000, "stock": 402,
        "slug": "serum-despigmentante-corporal-smooth-legend-90ml", "nota": "Tratamiento post-depilación que aclara folículos oscuros y previene vellos encarnados."
    },
    {
        "id": "083", "cat": "Cuidado Corporal", "sub": "Piernas, Pies y Manos",
        "name": "Crema Reparadora Intensiva con Urea al 40% para Talones y Codos",
        "supplier": "477266", "dropiId": "1681615", "costo": 16000, "pvp": 49000, "stock": 1162,
        "slug": "crema-reparadora-intensiva-urea-40-talones-codos", "nota": "Queratolítico clínico para reparar grietas profundas y callosidades en pies y codos."
    },
    {
        "id": "084", "cat": "Cuidado Corporal", "sub": "Piernas, Pies y Manos",
        "name": "Crema Alivio y Descanso para Piernas Cansadas Vital 100ml",
        "supplier": "29636", "dropiId": "2094216", "costo": 10000, "pvp": 39000, "stock": 5104,
        "slug": "crema-alivio-descanso-piernas-cansadas-vital-100ml", "nota": "Castaño de indias y mentol con efecto venotónico frío que alivia pesadez e hinchazón."
    },
    {
        "id": "085", "cat": "Cuidado Corporal", "sub": "Piernas, Pies y Manos",
        "name": "Mousse Efervescente para Piernas Cansadas con Efecto Hielo 150ml",
        "supplier": "263298", "dropiId": "959395", "costo": 30000, "pvp": 69000, "stock": 500,
        "slug": "mousse-efervescente-piernas-cansadas-efecto-hielo-150ml", "nota": "Espuma crepitante criogénica que intensifica la vasoconstricción y el retorno venoso."
    },
    {
        "id": "086", "cat": "Cuidado Corporal", "sub": "Piernas, Pies y Manos",
        "name": "Aceite Regenerador Puro de Rosa Mosqueta Chilena 30ml",
        "supplier": "336981", "dropiId": "1420003", "costo": 27350, "pvp": 62000, "stock": 499,
        "slug": "aceite-regenerador-puro-rosa-mosqueta-chilena-30ml", "nota": "100% puro prensado en frío con ácidos omega 3, 6 y 9 para cicatrices y estrías."
    },
    {
        "id": "087", "cat": "Cuidado Corporal", "sub": "Nutrición, Bronceado y Cera",
        "name": "Mantequilla Corporal Nutritiva Batida Truly con Karité 135ml",
        "supplier": "525739", "dropiId": "1851172", "costo": 16000, "pvp": 49000, "stock": 3648,
        "slug": "mantequilla-corporal-nutritiva-batida-truly-karite-135ml", "nota": "Body butter batida ultra emoliente que repara la barrera lipídica sin sensación grasa."
    },
    {
        "id": "088", "cat": "Cuidado Corporal", "sub": "Nutrición, Bronceado y Cera",
        "name": "Gotas Autobronceadoras Graduales Sun Drops Truly 90ml",
        "supplier": "525739", "dropiId": "2073347", "costo": 24000, "pvp": 65000, "stock": 2936,
        "slug": "gotas-autobronceadoras-graduales-sun-drops-truly-90ml", "nota": "DHA vegetal y ácido hialurónico para un bronceado natural gradual sin sol ni daño UV."
    },
    {
        "id": "089", "cat": "Cuidado Corporal", "sub": "Nutrición, Bronceado y Cera",
        "name": "Olla Calentadora Profesional de Cera Depilatoria Pro-Wax 100",
        "supplier": "48587", "dropiId": "263558", "costo": 30900, "pvp": 75000, "stock": 250,
        "slug": "olla-calentadora-cera-depilatoria-prowax-100", "nota": "Termostato regulable con cubeta de aluminio para depilación higiénica a temperatura segura."
    },
    {
        "id": "090", "cat": "Cuidado Corporal", "sub": "Nutrición, Bronceado y Cera",
        "name": "Gel Conductor Neutro Electrotópico para Cavitación y RF 500ml",
        "supplier": "365543", "dropiId": "1736707", "costo": 10000, "pvp": 35000, "stock": 1975,
        "slug": "gel-conductor-neutro-electrotopico-cavitacion-rf-500ml", "nota": "Acople acústico e impedancia neutra indispensable para cavitación, radiofrecuencia y EMS."
    }
]

def update_excel():
    print(f"Cargando libro: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Iniciar en la fila 62
    start_row = 62

    for i, p in enumerate(CORPORAL_PRODUCTS):
        row = start_row + i
        
        ws[f"A{row}"] = p["id"]
        ws[f"B{row}"] = p["cat"]
        ws[f"C{row}"] = p["sub"]
        ws[f"D{row}"] = p["name"]
        ws[f"E{row}"] = p["supplier"]
        ws[f"F{row}"] = p["dropiId"]
        ws[f"G{row}"] = p["costo"]
        ws[f"H{row}"] = p["pvp"]
        
        # Fórmulas de margen
        ws[f"I{row}"] = f"=H{row}-G{row}"
        ws[f"J{row}"] = f"=I{row}/H{row}"
        ws[f"K{row}"] = 16000  # Costo flete promedio COD
        ws[f"L{row}"] = f"=I{row}-K{row}"
        ws[f"M{row}"] = f"=L{row}/H{row}"
        
        ws[f"N{row}"] = p["stock"]
        ws[f"O{row}"] = "4.8 / Verificado"
        ws[f"P{row}"] = "Aprobado"
        ws[f"Q{row}"] = f"https://mesolabpro.com.co/tienda/corporal/{p['slug']}"
        ws[f"R{row}"] = p["nota"]

        # Formatos numéricos
        ws[f"G{row}"].number_format = "$#,##0"
        ws[f"H{row}"].number_format = "$#,##0"
        ws[f"I{row}"].number_format = "$#,##0"
        ws[f"J{row}"].number_format = "0.0%"
        ws[f"K{row}"].number_format = "$#,##0"
        ws[f"L{row}"].number_format = "$#,##0"
        ws[f"M{row}"].number_format = "0.0%"
        ws[f"N{row}"].number_format = "#,##0"

    wb.save(EXCEL_PATH)
    print(f"✓ Éxito: 30 productos de Cuidado Corporal insertados en filas {start_row} a {start_row + len(CORPORAL_PRODUCTS) - 1}.")

if __name__ == "__main__":
    update_excel()
